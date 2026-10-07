"""Load workflow definitions from the Excel workbook (the single source of truth).

Sheets
  Workflows       (provided)  - id, name, trigger, inputs, steps, decision logic, tools, output
  Test_Questions  (provided)  - example requests per workflow
  Input_Spec      (added)     - machine-readable inputs for each workflow
  Step_Bindings   (added)     - binds every step in the Steps column to a registered tool
  Routing_Hints   (added)     - extra vocabulary that helps the router
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from .errors import WorkflowConfigError
from .models import InputSpec, StepDef, WorkflowDef


def _rows(ws):
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    for i, r in enumerate(it, start=2):
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        yield i, {header[j]: r[j] for j in range(len(header)) if header[j]}


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _bool(v) -> bool:
    return _s(v).lower() in {"y", "yes", "true", "1"}


def load_workflows(path: Path) -> dict[str, WorkflowDef]:
    path = Path(path)
    if not path.exists():
        raise WorkflowConfigError(f"Workflow workbook not found: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in ("Workflows", "Input_Spec", "Step_Bindings"):
        if sheet not in wb.sheetnames:
            raise WorkflowConfigError(f"Workbook is missing required sheet '{sheet}'")

    wfs: dict[str, WorkflowDef] = {}
    for _, r in _rows(wb["Workflows"]):
        wid = _s(r.get("Workflow_ID"))
        wfs[wid] = WorkflowDef(
            id=wid, name=_s(r.get("Workflow_Name")), trigger=_s(r.get("Trigger")),
            inputs_text=_s(r.get("Inputs")), steps_text=_s(r.get("Steps")),
            decision_logic=_s(r.get("Decision_Logic")), tools_text=_s(r.get("Tools_Required")),
            expected_output=_s(r.get("Expected_Output")),
        )

    if "Test_Questions" in wb.sheetnames:
        for _, r in _rows(wb["Test_Questions"]):
            wid = _s(r.get("Workflow_ID"))
            if wid in wfs and _s(r.get("Test_Request")):
                wfs[wid].examples.append(_s(r.get("Test_Request")))

    if "Routing_Hints" in wb.sheetnames:
        for _, r in _rows(wb["Routing_Hints"]):
            wid = _s(r.get("Workflow_ID"))
            if wid in wfs:
                wfs[wid].keywords = _s(r.get("Keywords"))

    for ln, r in _rows(wb["Input_Spec"]):
        wid = _s(r.get("Workflow_ID"))
        if wid not in wfs:
            raise WorkflowConfigError(f"Input_Spec row {ln}: unknown workflow '{wid}'")
        default = r.get("Default")
        wfs[wid].input_specs.append(InputSpec(
            name=_s(r.get("Input_Name")), type=_s(r.get("Type")) or "text",
            required=_bool(r.get("Required")), default=None if _s(default) == "" else default,
            description=_s(r.get("Description")), ask=_s(r.get("Ask_Prompt")),
            extract_regex=_s(r.get("Extract_Regex")),
            fields=[f.strip() for f in _s(r.get("Fields")).split(",") if f.strip()],
        ))

    for ln, r in _rows(wb["Step_Bindings"]):
        wid = _s(r.get("Workflow_ID"))
        if wid not in wfs:
            raise WorkflowConfigError(f"Step_Bindings row {ln}: unknown workflow '{wid}'")
        raw = _s(r.get("Args_JSON")) or "{}"
        try:
            args = json.loads(raw)
        except json.JSONDecodeError as e:
            raise WorkflowConfigError(f"Step_Bindings row {ln} ({wid}): Args_JSON is not valid JSON: {e}")
        wfs[wid].steps.append(StepDef(
            no=int(r.get("Step_No")), name=_s(r.get("Step_Name")), tool=_s(r.get("Tool")),
            args=args, output_key=_s(r.get("Output_Key")) or f"step{r.get('Step_No')}",
            condition=_s(r.get("Condition")), on_error=_s(r.get("On_Error")) or "abort",
            note=_s(r.get("Note")),
        ))
    for wf in wfs.values():
        wf.steps.sort(key=lambda s: s.no)
    return wfs


def validate_workflows(wfs: dict[str, WorkflowDef]) -> list[str]:
    """Static checks so a bad row in Excel is caught at load time, not mid-run."""
    import re
    from .registry import list_tools
    problems = []
    tools = list_tools()
    tpl = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")

    def templates(o):
        if isinstance(o, str):
            yield from tpl.findall(o)
        elif isinstance(o, dict):
            for v in o.values():
                yield from templates(v)
        elif isinstance(o, list):
            for v in o:
                yield from templates(v)

    for wf in wfs.values():
        if not wf.steps:
            problems.append(f"{wf.id}: no steps defined in Step_Bindings")
        names = {s.name for s in wf.input_specs}
        produced = set()
        nos = [s.no for s in wf.steps]
        if nos != list(range(1, len(nos) + 1)):
            problems.append(f"{wf.id}: step numbers should be 1..N without gaps, got {nos}")
        for s in wf.steps:
            t = tools.get(s.tool)
            if not t:
                problems.append(f"{wf.id} step {s.no}: unknown tool '{s.tool}'")
            else:
                for a in t.required_args:
                    if a not in s.args:
                        problems.append(f"{wf.id} step {s.no}: tool '{s.tool}' needs arg '{a}'")
            for ref in templates(s.args):
                root = ref.split(".")[0].split("[")[0]
                if root == "inputs":
                    key = ref.split(".")[1] if "." in ref else ""
                    if key not in names:
                        problems.append(f"{wf.id} step {s.no}: '{{{{{ref}}}}}' not declared in Input_Spec")
                elif root == "ctx":
                    key = ref.split(".")[1].split("[")[0] if "." in ref else ""
                    if key not in produced:
                        problems.append(f"{wf.id} step {s.no}: '{{{{{ref}}}}}' used before a step produces it")
                elif root not in ("run", "workflow", "config"):
                    problems.append(f"{wf.id} step {s.no}: unknown template root in '{{{{{ref}}}}}'")
            produced.add(s.output_key)
        if "result" not in produced:
            problems.append(f"{wf.id}: no step writes Output_Key 'result' (the final answer)")
    return problems
