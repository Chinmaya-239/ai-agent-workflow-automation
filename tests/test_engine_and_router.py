import json
import shutil
from pathlib import Path

import openpyxl
import pytest

from agent.errors import ToolError, WorkflowConfigError
from agent.engine import parse_policy
from agent.excel_loader import validate_workflows, load_workflows
from agent.llm.base import LLMClient
from agent.errors import LLMError
from agent.registry import tool
from agent.safe_eval import safe_eval
from agent.templating import resolve

ROOT = Path(__file__).resolve().parent.parent


# ---------------------------------------------------------------- safe_eval / templating
def test_safe_eval_allows_rules_blocks_code():
    assert safe_eval("a < coalesce(b, c)", {"a": 1, "b": None, "c": 5}) is True
    assert safe_eval("x['count'] > 0 and not y", {"x": {"count": 2}, "y": False})
    for evil in ("__import__('os').system('echo hi')", "open('/etc/passwd')", "().__class__", "exec('1')"):
        with pytest.raises(WorkflowConfigError):
            safe_eval(evil, {})


def test_templating_native_types_and_errors():
    ns = {"ctx": {"a": {"rows": [1, 2]}}, "inputs": {"n": 5}}
    assert resolve("{{ctx.a.rows}}", ns) == [1, 2]
    assert resolve("n={{inputs.n}}!", ns) == "n=5!"
    assert resolve({"k": ["{{ctx.a.rows[1]}}"]}, ns) == {"k": [2]}
    with pytest.raises(WorkflowConfigError):
        resolve("{{ctx.nope}}", ns)


def test_error_policy_parser():
    assert parse_policy("retry:2+skip") == (2, "skip") and parse_policy("abort") == (0, "abort")
    with pytest.raises(WorkflowConfigError):
        parse_policy("explode")


# ---------------------------------------------------------------- failure handling
def test_missing_file_fails_gracefully(agent):
    r = agent.run("which products need restocking", {"inventory_path": "data/does_not_exist.csv"})
    assert r.status == "failed" and "File not found" in r.error and r.steps[0].status == "failed"


def test_missing_column_fails_with_clear_message(agent, tmp_path):
    p = tmp_path / "inv.csv"; p.write_text("sku,name\nA,B\n")
    r = agent.run("which products need restocking", {"inventory_path": str(p)})
    assert r.status == "failed" and "Missing required column" in r.error


def test_unexpected_tool_exception_does_not_crash_agent(agent):
    @tool("boom_test_tool")
    def boom(args, tc):
        raise RuntimeError("kaboom")
    wf = agent.workflows["WF001"]
    from agent.models import StepDef
    wf.steps.insert(0, StepDef(0, "boom", "boom_test_tool", {}, "boom"))
    r = agent.run("which products need restocking")
    wf.steps.pop(0)
    assert r.status == "failed" and "RuntimeError" in r.error


# ---------------------------------------------------------------- workbook validation
def test_workbook_lint_catches_bad_rows(tmp_path):
    p = tmp_path / "bad.xlsx"
    shutil.copy(ROOT / "workflows" / "ai_agent_workflows.xlsx", p)
    wb = openpyxl.load_workbook(p)
    ws = wb["Step_Bindings"]
    ws["D2"] = "no_such_tool"
    ws["E3"] = json.dumps({"rows": "{{ctx.never_made.rows}}", "where": "x"})
    wb.save(p)
    problems = validate_workflows(load_workflows(p))
    assert any("unknown tool 'no_such_tool'" in x for x in problems)
    assert any("used before a step produces it" in x for x in problems)


def test_shipped_workbook_is_clean():
    assert validate_workflows(load_workflows(ROOT / "workflows" / "ai_agent_workflows.xlsx")) == []


# ---------------------------------------------------------------- routing
TEST_QUESTIONS = {
    "Which products need restocking?": "WF001", "Find products where vendor price differs by more than 10%.": "WF002",
    "Process this vendor spreadsheet and show invalid rows.": "WF003", "Generate SEO content for this product.": "WF004",
    "Where is order ORD-1001?": "WF005", "Find likely duplicate products in the catalog.": "WF006",
    "Create a campaign brief for the new collection.": "WF007", "Classify these keywords and map them to pages.": "WF008",
    "Assign this urgent task to the best available developer.": "WF009", "Which workflows are failing most often?": "WF010",
    "Anything running low in the warehouse?": "WF001", "has alice@example.com's package shipped yet?": "WF005",
    "are there any repeated listings in our catalog?": "WF006", "who should take this bug fix?": "WF009",
}


@pytest.mark.parametrize("q,wf", TEST_QUESTIONS.items())
def test_router_offline(agent, q, wf):
    assert agent.router.route(q).workflow_id == wf


def test_router_refuses_to_guess(agent):
    r = agent.run("What's the weather in Hyderabad?")
    assert r.status == "clarify" and not r.steps and "Did you mean" in r.message


# ---------------------------------------------------------------- LLM paths via a fake provider
class FakeLLM(LLMClient):
    name = "fake"

    def __init__(self, route="WF001", fail=False):
        self.route, self.fail, self.calls = route, fail, []

    def structured(self, system, user, tool_name, tool_description, schema):
        self.calls.append(tool_name)
        if self.fail:
            raise LLMError("simulated outage")
        if tool_name == "select_workflow":
            return {"workflow_id": self.route, "confidence": 0.9, "reasoning": "fake reasoning"}
        if tool_name == "extract_inputs":
            return {}
        if tool_name == "classify_keywords":
            n = user.count("\n")
            return {"items": [{"index": i, "intent": "informational"} for i in range(n)]}
        raise LLMError("unsupported in fake")

    def text(self, system, user, max_tokens=700):
        self.calls.append("text")
        if self.fail:
            raise LLMError("simulated outage")
        return "LLM copy"


def test_llm_router_tool_call_used(make_agent):
    llm = FakeLLM(route="WF006")
    a = make_agent(llm)
    r = a.run("anything at all about the catalog")
    assert r.workflow_id == "WF006" and r.routing.method == "llm-tool-call" and "select_workflow" in llm.calls


def test_llm_returns_unknown_workflow_falls_back_to_keywords(make_agent):
    a = make_agent(FakeLLM(route="WF999"))
    r = a.run("Which products need restocking?")
    assert r.workflow_id == "WF001" and r.routing.method == "keyword-tfidf" and "LLM routing failed" in r.routing.reasoning


def test_llm_outage_everywhere_still_completes(make_agent):
    a = make_agent(FakeLLM(fail=True))
    r = a.run("Classify these keywords and map them to pages.")
    assert r.status == "success"
    r = a.run("generate content", {"product_attributes": "name: Mug; color: red"}, "WF004")
    assert r.status == "success" and any("template fallback" in n for s in r.steps for n in s.notes)


def test_llm_text_used_when_available(make_agent):
    a = make_agent(FakeLLM())
    r = a.run("x", {"product_attributes": "name: Mug"}, "WF004")
    secs = {s["heading"]: s["content"] for s in r.result["sections"]}
    assert secs["Product description"] == "LLM copy"


def test_llm_classification_with_unknown_items_uses_rules(make_agent):
    a = make_agent(FakeLLM())
    r = a.run("Classify these keywords", workflow_id="WF008")
    assert r.status == "success"


# ---------------------------------------------------------------- 11th workflow, zero Python
def test_eleventh_workflow_needs_no_code(make_agent):
    import subprocess, sys
    subprocess.run([sys.executable, str(ROOT / "scripts" / "add_wf011_example.py")], check=True, capture_output=True)
    a = make_agent(workbook=ROOT / "workflows" / "ai_agent_workflows_with_wf011.xlsx")
    assert len(a.workflows) == 11
    r = a.run("Which products have a low margin?")
    assert r.workflow_id == "WF011" and r.status == "success"
    r = a.run("Which products need restocking?")          # old workflows are unaffected
    assert r.workflow_id == "WF001"
