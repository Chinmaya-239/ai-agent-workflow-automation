"""Build workflows/ai_agent_workflows.xlsx = the ORIGINAL assessment sheets (untouched)
+ three machine-readable sheets that make every workflow executable:

  Input_Spec      what each workflow needs (type, required?, default, how to ask the user, extraction regex)
  Step_Bindings   which registered tool runs each step, with templated args, condition and error policy
  Routing_Hints   extra vocabulary for the router

Adding an 11th workflow = add rows to Workflows + these three sheets. No Python needed unless a brand-new tool is required.
Run: python scripts/build_workbook.py
"""
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
SRC, OUT = ROOT / "workflows" / "original_assessment.xlsx", ROOT / "workflows" / "ai_agent_workflows.xlsx"

J = lambda **kw: json.dumps(kw, ensure_ascii=False)

# ---------------------------------------------------------------- Routing_Hints
HINTS = {
    "WF001": "restock restocking inventory stock low stock reorder running out out of stock replenish warehouse",
    "WF002": "price prices pricing validate validation vendor price price list differs difference mismatch discrepancy compare prices",
    "WF003": "vendor file spreadsheet process clean cleanup invalid rows normalize columns supplier file upload ingest",
    "WF004": "generate content product description seo title meta description short description copy write listing",
    "WF005": "order status where is order track tracking shipment shipping delivery ORD customer order package",
    "WF006": "duplicate duplicates duplicated catalog likely same product dedupe repeated listings",
    "WF007": "campaign brief marketing campaign promotion plan collection launch channels",
    "WF008": "keyword keywords classify classification search intent seo keywords map pages target page",
    "WF009": "assign task employee developer best available workload skills delegate who should",
    "WF010": "performance report workflows failing failure rate execution time slow errors logs success rate",
}

# ---------------------------------------------------------------- Input_Spec
# (wf, name, type, required, default, description, ask_prompt, extract_regex, fields)
INPUTS = [
    ("WF001", "inventory_path", "file", "N", "data/inventory.csv", "Product inventory CSV", "", "", ""),
    ("WF001", "threshold_override", "number", "N", "", "Optional global minimum stock threshold (overrides per-product minimum)", "", r"threshold\s*(?:of|=|:|is)?\s*(\d+(?:\.\d+)?)", ""),
    ("WF002", "products_path", "file", "N", "data/products.csv", "Internal product CSV (sku, product_name, price)", "", "", ""),
    ("WF002", "vendor_prices_path", "file", "N", "data/vendor_prices.csv", "Vendor price list CSV (sku, vendor_price)", "", "", ""),
    ("WF002", "threshold_pct", "number", "N", "10", "Price difference % above which a product is flagged", "", r"(\d+(?:\.\d+)?)\s*%", ""),
    ("WF003", "vendor_file", "file", "Y", "data/vendor_file_sample.csv", "Vendor CSV/XLSX file to process", "Please provide the vendor file (.csv or .xlsx) to process.", "", ""),
    ("WF004", "product_attributes", "object", "Y", "", "Product details provided by the user", "Please describe the product (name, category, material, color, target audience).", "", "product_name,category,material,color,target_audience,features"),
    ("WF005", "identifier", "identifier", "Y", "", "Order ID (ORD-1234) or customer email", "Please provide the order ID (e.g. ORD-1001) or the customer email.", "", ""),
    ("WF005", "orders_path", "file", "N", "data/orders.json", "Order database (simulated)", "", "", ""),
    ("WF005", "shipments_path", "file", "N", "data/shipments.json", "Shipment/carrier API (simulated)", "", "", ""),
    ("WF006", "catalog_path", "file", "N", "data/catalog.csv", "Product catalog CSV", "", "", ""),
    ("WF006", "similarity_threshold", "number", "N", "0.75", "Minimum similarity (0-1) for a *possible* duplicate", "", r"similarity\s*(?:of|=|:|>=?)?\s*(0?\.\d+)", ""),
    ("WF007", "campaign_goal", "text", "Y", "", "Campaign goal", "What is the goal of the campaign (e.g. 'increase online sales by 20%')?", r"goal\s*(?:is|:)\s*(.+?)(?=\.\s|\.$|;|\n|$)", ""),
    ("WF007", "campaign_dates", "date_range", "Y", "", "Campaign start and end dates", "What are the campaign start and end dates (e.g. 2026-11-01 to 2026-11-30)?", "", ""),
    ("WF007", "target_audience", "text", "N", "", "Target audience", "", r"audience\s*(?:is|:)\s*(.+?)(?=\.\s|\.$|;|\n|$)", ""),
    ("WF007", "promotion", "text", "N", "", "Promotion / offer", "", r"promotion\s*(?:is|:)\s*(.+?)(?=\.\s|\.$|;|\n|$)", ""),
    ("WF007", "product_filter", "text", "N", "", "Which products the campaign is about", "", r"\bfor\s+(?:the\s+)?(.+?(?:collection|range|products|line))", ""),
    ("WF007", "products_path", "file", "N", "data/products.csv", "Product data CSV", "", "", ""),
    ("WF008", "keywords_path", "file", "Y", "data/keywords.csv", "Keyword CSV (keyword, search_volume, difficulty)", "Please provide a keyword CSV.", "", ""),
    ("WF008", "categories_path", "file", "N", "data/seo_categories.csv", "Category/page map (category, page_url, terms)", "", "", ""),
    ("WF009", "task_description", "text", "Y", "", "Task to assign", "Please describe the task to assign.", "REQUEST", ""),
    ("WF009", "employees_path", "file", "N", "data/employees.csv", "Employee, skills and workload data", "", "", ""),
    ("WF009", "priority", "text", "N", "", "Explicit priority (high/medium/low); otherwise inferred from the text", "", "", ""),
    ("WF009", "deadline", "text", "N", "", "Explicit deadline; otherwise parsed from the text or assumed from priority", "", "", ""),
    ("WF009", "effort_hours", "number", "N", "", "Estimated effort in hours; otherwise parsed or assumed 8h", "", "", ""),
    ("WF010", "log_files", "file_list", "N", "data/execution_logs.csv;logs/runtime_logs.csv", "Execution log CSV files", "", "", ""),
    ("WF010", "failure_rate_threshold_pct", "number", "N", "10", "Flag workflows whose failure rate exceeds this %", "", "", ""),
    ("WF010", "exec_time_threshold_sec", "number", "N", "20", "Flag workflows whose average run time exceeds this many seconds", "", "", ""),
    ("WF010", "slow_step_threshold_sec", "number", "N", "5", "A step is 'slow' when its average duration exceeds this", "", "", ""),
]

# ---------------------------------------------------------------- Step_Bindings
# (wf, no, name, tool, args(dict), output_key, condition, on_error, note)
S = []
def step(wf, no, name, tool, args, out, cond="", on_error="abort", note=""):
    S.append((wf, no, name, tool, json.dumps(args, ensure_ascii=False), out, cond, on_error, note))

# ---- WF001
step("WF001", 1, "Load inventory", "load_table", dict(path="{{inputs.inventory_path}}", required_columns=["sku", "product_name", "current_stock", "minimum_stock"], numeric_columns=["current_stock", "minimum_stock"], optional_numeric_columns=["reorder_target"]), "inventory", note="Rows with non-numeric stock are quarantined, not silently dropped")
step("WF001", 2, "Compare current stock with minimum threshold", "compute_columns", dict(rows="{{ctx.inventory.rows}}", columns={"threshold": "coalesce(threshold_override, minimum_stock)", "below_threshold": "current_stock < threshold"}), "compared", note="Decision rule: stock strictly below threshold")
step("WF001", 3, "Identify low-stock products", "filter_rows", dict(rows="{{ctx.compared.rows}}", where="below_threshold", rule="current_stock < minimum_stock (threshold)"), "low_stock", note="Excel rule: If current_stock < minimum_stock, mark product for restock")
step("WF001", 4, "Calculate reorder quantity", "compute_columns", dict(rows="{{ctx.low_stock.rows}}", columns={"target_stock": "coalesce(reorder_target, threshold * 2)", "suggested_reorder_qty": "max(target_stock - current_stock, 0)", "shortfall": "threshold - current_stock"}), "restock_rows", note="Reorder up to reorder_target (or 2x threshold when no target is given)")
step("WF001", 5, "Generate restock list", "make_report", dict(title="Restock list", summary="{{ctx.low_stock.count}} of {{ctx.low_stock.total}} products are below their minimum stock and need restocking.", metrics={"Products checked": "{{ctx.inventory.row_count}}", "Rows skipped (bad data)": "{{ctx.inventory.invalid_count}}"}, tables={"Products requiring restock": {"rows": "{{ctx.restock_rows.rows}}", "columns": ["sku", "product_name", "current_stock", "threshold", "suggested_reorder_qty"], "sort_by": "shortfall", "descending": True, "export": True}, "Rows skipped (invalid data)": {"rows": "{{ctx.inventory.invalid_rows}}", "columns": ["row_number", "sku", "invalid_reason"]}}), "result")

# ---- WF002
step("WF002", 1, "Load product prices", "load_table", dict(path="{{inputs.products_path}}", required_columns=["sku", "product_name", "price"], numeric_columns=["price"]), "products")
step("WF002", 2, "Load vendor price list", "load_table", dict(path="{{inputs.vendor_prices_path}}", required_columns=["sku", "vendor_price"], numeric_columns=["vendor_price"]), "vendor")
step("WF002", 3, "Match products by SKU", "join_tables", dict(left="{{ctx.products.rows}}", right="{{ctx.vendor.rows}}", key="sku", bring=["vendor_price"]), "matched", note="SKUs compared after upper-casing and removing spaces/dashes")
step("WF002", 4, "Compare internal and vendor prices", "compute_columns", dict(rows="{{ctx.matched.rows}}", columns={"price_diff": "round(price - vendor_price, 2)"}), "compared")
step("WF002", 5, "Calculate percentage difference", "compute_columns", dict(rows="{{ctx.compared.rows}}", columns={"pct_diff": "round((price - vendor_price) / vendor_price * 100, 2) if vendor_price else None", "abs_pct_diff": "abs(pct_diff) if pct_diff is not None else None", "status": "'CANNOT VALIDATE' if pct_diff is None else ('EXCEPTION' if abs(pct_diff) > threshold_pct else 'OK')"}), "scored", note="% difference is relative to the vendor price")
step("WF002", 6, "Flag exceptions", "filter_rows", dict(rows="{{ctx.scored.rows}}", where="status != 'OK'", rule="|price difference %| > threshold_pct (or cannot be computed)"), "exceptions", note="Excel rule: flag when difference exceeds 10% (strictly greater)")
step("WF002", 7, "Generate validation report", "make_report", dict(title="Product price validation", summary="{{ctx.matched.count}} products matched by SKU; {{ctx.exceptions.count}} exception(s) where the price difference exceeds {{inputs.threshold_pct}}% (or cannot be validated).", metrics={"Internal products": "{{ctx.products.row_count}}", "Matched": "{{ctx.matched.count}}", "No vendor price": "{{ctx.matched.unmatched_left_count}}", "Vendor-only SKUs": "{{ctx.matched.unmatched_right_count}}"}, tables={"Exceptions": {"rows": "{{ctx.exceptions.rows}}", "columns": ["sku", "product_name", "price", "vendor_price", "price_diff", "pct_diff", "status"], "sort_by": "abs_pct_diff", "descending": True, "export": True}, "All matched products": {"rows": "{{ctx.scored.rows}}", "columns": ["sku", "price", "vendor_price", "pct_diff", "status"], "export": True, "limit": 15}, "Internal SKUs missing from vendor list": {"rows": "{{ctx.matched.unmatched_left}}", "columns": ["sku", "product_name", "price"]}, "Vendor SKUs not in our catalog": {"rows": "{{ctx.matched.unmatched_right}}", "columns": ["sku", "vendor_price"]}}), "result")

# ---- WF003
step("WF003", 1, "Read file", "load_table", dict(path="{{inputs.vendor_file}}", normalize_headers=False), "raw")
step("WF003", 2, "Detect columns", "detect_columns", dict(columns="{{ctx.raw.columns}}", required=["sku", "product_name"]), "colmap", note="Exact synonym > keyword > fuzzy match")
step("WF003", 3, "Normalize column names", "normalize_vendor_rows", dict(rows="{{ctx.raw.rows}}", mapping="{{ctx.colmap.mapping}}"), "normalized")
step("WF003", 4, "Validate required fields", "validate_rows", dict(rows="{{ctx.normalized.rows}}", required=["sku", "product_name"]), "validated", note="Excel rule: rows missing SKU or product name are invalid")
step("WF003", 5, "Identify invalid rows", "partition_rows", dict(rows="{{ctx.validated.rows}}", flag="is_valid"), "split")
step("WF003", 6, "Produce cleaned dataset", "export_tables", dict(tables={"cleaned": "{{ctx.split.valid_rows}}", "invalid_rows": "{{ctx.split.invalid_rows}}"}, directory="{{config.outputs_dir}}", prefix="wf003", columns={"invalid_rows": ["row_number", "invalid_reason", "sku", "product_name", "price", "quantity", "category"]}), "export")
step("WF003", 7, "Generate validation summary", "make_report", dict(title="Vendor file processing", summary="{{ctx.split.valid_count}} of {{ctx.split.total}} rows are valid; {{ctx.split.invalid_count}} invalid row(s) were separated.", metrics={"Source file": "{{inputs.vendor_file}}", "Detected columns": "{{ctx.colmap.mapping}}"}, tables={"Invalid rows": {"rows": "{{ctx.split.invalid_rows}}", "columns": ["row_number", "invalid_reason", "sku", "product_name", "price"]}, "Cleaned dataset (preview)": {"rows": "{{ctx.split.valid_rows}}", "columns": ["sku", "product_name", "price", "quantity", "category", "color"], "limit": 8}}, files="{{ctx.export.files}}"), "result")

# ---- WF004
ATTR = "{{ctx.attrs}}"
step("WF004", 1, "Validate required attributes", "validate_attributes", dict(attributes="{{inputs.product_attributes}}", required=["product_name"], expected=["category", "material", "color", "target_audience", "features"]), "attrs", note="Excel rule: do not invent missing attributes - mark them as missing")
step("WF004", 2, "Create product description", "generate_text", dict(field="description", attributes=ATTR, instructions="Write a persuasive product description of 50-80 words.", fallback_template="The {product_name}<< ({category})>><< made from {material}>><< in {color}>>.<< It is designed for {target_audience}.>><< Key features: {features}.>>"), "description", on_error="retry:1")
step("WF004", 3, "Generate short description", "generate_text", dict(field="short_description", attributes=ATTR, instructions="Write a short description of at most 25 words.", fallback_template="{product_name}<< - {category}>><< in {color}>><< for {target_audience}>>."), "short_description", on_error="retry:1")
step("WF004", 4, "Generate SEO title", "generate_text", dict(field="seo_title", attributes=ATTR, instructions="Write an SEO page title of at most 60 characters. No quotes.", fallback_template="{product_name}<< - {category}>><< ({color})>>"), "seo_title", on_error="retry:1")
step("WF004", 5, "Generate meta description", "generate_text", dict(field="meta_description", attributes=ATTR, instructions="Write a meta description of at most 155 characters with a soft call to action.", fallback_template="Shop the {product_name}<<, a {material} {category}>><< in {color}>><< for {target_audience}>>. Order online today."), "meta_description", on_error="retry:1")
step("WF004", 6, "Validate generated text", "validate_text", dict(texts={"description": "{{ctx.description}}", "short_description": "{{ctx.short_description}}", "seo_title": "{{ctx.seo_title}}", "meta_description": "{{ctx.meta_description}}"}, limits={"short_description": 160, "seo_title": 60, "meta_description": 155}, attributes=ATTR), "checked")
step("WF004", 7, "Compile product content", "make_report", dict(title="Product content", summary="Content generated using only the attributes you provided.", sections={"Product description": "{{ctx.checked.texts.description}}", "Short description": "{{ctx.checked.texts.short_description}}", "SEO title": "{{ctx.checked.texts.seo_title}}", "Meta description": "{{ctx.checked.texts.meta_description}}", "Attributes used": "{{ctx.attrs.provided}}", "MISSING information (not invented)": "{{ctx.attrs.missing}}", "Validation warnings": "{{ctx.checked.warnings}}"}), "result")

# ---- WF005
step("WF005", 1, "Validate identifier", "validate_identifier", dict(value="{{inputs.identifier}}"), "ident")
step("WF005", 2, "Search order data", "find_orders", dict(identifier="{{ctx.ident}}", source="{{inputs.orders_path}}"), "orders", note="Excel rule: if no order is found, ask for another identifier")
step("WF005", 3, "Retrieve order status", "get_order_status", dict(orders="{{ctx.orders.orders}}"), "statuses")
step("WF005", 4, "Retrieve shipment information", "get_shipment_info", dict(orders="{{ctx.orders.orders}}", source="{{inputs.shipments_path}}"), "shipments", cond="orders['with_shipment_count'] > 0", on_error="retry:2+skip", note="Carrier API may time out: retry twice, then continue without tracking")
step("WF005", 5, "Summarize current status", "summarize_orders", dict(statuses="{{ctx.statuses}}", shipments="{{ctx.shipments}}"), "result")

# ---- WF006
step("WF006", 1, "Load products", "load_table", dict(path="{{inputs.catalog_path}}", required_columns=["sku", "product_name"]), "catalog")
step("WF006", 2, "Normalize names/SKUs", "normalize_fields", dict(rows="{{ctx.catalog.rows}}"), "normalized")
step("WF006", 3, "Compare identifiers", "group_exact", dict(rows="{{ctx.normalized.rows}}", key="norm_sku"), "exact", note="Excel rule: exact SKU match is a definite duplicate")
step("WF006", 4, "Compare product attributes", "find_similar_pairs", dict(rows="{{ctx.normalized.rows}}", attribute_fields=["brand", "category", "color", "size"], block_field="category", threshold="{{inputs.similarity_threshold}}", exclude_groups="{{ctx.exact.groups}}"), "similar", note="Excel rule: high attribute similarity is a possible duplicate")
step("WF006", 5, "Group likely duplicates", "build_duplicate_groups", dict(rows="{{ctx.normalized.rows}}", exact="{{ctx.exact}}", similar="{{ctx.similar}}"), "groups")
step("WF006", 6, "Assign confidence", "assign_confidence", dict(groups="{{ctx.groups.groups}}", rows="{{ctx.groups.rows}}", high_threshold=0.9), "scored")
step("WF006", 7, "Generate duplicate report", "make_report", dict(title="Duplicate product detection", summary="{{ctx.scored.definite}} definite and {{ctx.scored.possible}} possible duplicate group(s) found in {{ctx.catalog.row_count}} products.", tables={"Duplicate groups": {"rows": "{{ctx.scored.rows}}", "columns": ["group", "confidence", "score", "members", "example_name", "matching_fields", "differing_fields"], "export": True}}, notes=["Same product in a different colour or size is treated as a variant, not a duplicate."]), "result")

# ---- WF007
step("WF007", 1, "Validate inputs", "validate_campaign_inputs", dict(goal="{{inputs.campaign_goal}}", dates="{{inputs.campaign_dates}}", audience="{{inputs.target_audience}}", promotion="{{inputs.promotion}}"), "campaign", note="Excel rule: missing goal or dates are requested BEFORE generating (also enforced pre-flight via Input_Spec)")
step("WF007", 2, "Identify campaign objective", "classify_objective", dict(goal="{{ctx.campaign.goal}}"), "objective")
step("WF007", 3, "Load product data", "load_table", dict(path="{{inputs.products_path}}", required_columns=["sku", "product_name", "category", "price"], numeric_columns=["price"]), "products")
step("WF007", 4, "Summarize products", "summarize_products", dict(rows="{{ctx.products.rows}}", product_filter="{{inputs.product_filter}}"), "product_summary")
step("WF007", 5, "Create messaging", "create_messaging", dict(objective="{{ctx.objective}}", products="{{ctx.product_summary}}", audience="{{ctx.campaign.audience}}", promotion="{{ctx.campaign.promotion}}"), "messaging", on_error="retry:1")
step("WF007", 6, "Create channel recommendations", "recommend_channels", dict(objective="{{ctx.objective}}", audience="{{ctx.campaign.audience}}"), "channels")
step("WF007", 7, "Create campaign checklist", "build_timeline_checklist", dict(start="{{ctx.campaign.start}}", end="{{ctx.campaign.end}}", promotion="{{ctx.campaign.promotion}}", channels="{{ctx.channels}}"), "plan")
step("WF007", 8, "Generate campaign brief", "make_report", dict(title="Campaign brief", summary="{{ctx.messaging.headline}}", sections={"Objective": {"Goal": "{{ctx.objective.statement}}", "Type": "{{ctx.objective.label}}", "KPIs": "{{ctx.objective.kpis}}", "Target": "{{ctx.objective.target}}"}, "Audience": "{{ctx.campaign.audience}}", "Promotion": "{{ctx.campaign.promotion}}", "Dates": "{{ctx.campaign.start}} to {{ctx.campaign.end}} ({{ctx.campaign.duration_days}} days)", "Products": "{{ctx.product_summary.count}} products ({{ctx.product_summary.basis}}); categories {{ctx.product_summary.categories}}; price range {{ctx.product_summary.price_range}}", "Key messaging": "{{ctx.messaging.key_messages}}", "Call to action": "{{ctx.messaging.call_to_action}}", "Tone": "{{ctx.messaging.tone}}"}, tables={"Channels": {"rows": "{{ctx.channels.channels}}"}, "Timeline": {"rows": "{{ctx.plan.timeline}}"}, "Campaign checklist": {"rows": "{{ctx.plan.checklist}}", "export": True}}, notes=["Fields shown as 'Not specified' were not provided and have not been invented."]), "result")

# ---- WF008
step("WF008", 1, "Read keywords", "load_table", dict(path="{{inputs.keywords_path}}", required_columns=["keyword"], optional_numeric_columns=["search_volume", "difficulty"]), "keywords")
step("WF008", 2, "Remove duplicates", "dedupe_rows", dict(rows="{{ctx.keywords.rows}}", key="keyword"), "deduped")
step("WF008", 3, "Classify search intent", "classify_intent", dict(rows="{{ctx.deduped.rows}}", field="keyword", brand_terms=["acme"]), "classified", on_error="retry:1", note="LLM tool call with rule-based fallback; 4 allowed labels")
step("WF008", 4, "Map keywords to categories", "map_keywords_to_categories", dict(rows="{{ctx.classified.rows}}", categories_path="{{inputs.categories_path}}"), "mapped")
step("WF008", 5, "Identify high-priority keywords", "prioritize_keywords", dict(rows="{{ctx.mapped.rows}}", high_cutoff=55, medium_cutoff=35), "prioritized")
step("WF008", 6, "Export results", "export_tables", dict(tables={"keyword_report": "{{ctx.prioritized.rows}}"}, directory="{{config.outputs_dir}}", prefix="wf008", columns={"keyword_report": ["keyword", "search_volume", "difficulty", "intent", "category", "target_page", "priority_score", "priority"]}), "export")
step("WF008", 7, "Generate keyword report", "make_report", dict(title="SEO keyword report", summary="{{ctx.deduped.count}} unique keywords ({{ctx.deduped.removed_count}} duplicates removed); {{ctx.prioritized.high_count}} high priority; {{ctx.mapped.unmapped}} without a matching page.", metrics={"Intent distribution": "{{ctx.classified.intent_counts}}"}, tables={"Top priority keywords": {"rows": "{{ctx.prioritized.rows}}", "columns": ["keyword", "search_volume", "intent", "category", "target_page", "priority_score", "priority"], "limit": 12}, "Duplicates removed": {"rows": "{{ctx.deduped.removed}}", "columns": ["keyword", "duplicate_of"]}}, files="{{ctx.export.files}}"), "result")

# ---- WF009
step("WF009", 1, "Understand task requirements", "analyze_task", dict(description="{{inputs.task_description}}", employees_path="{{inputs.employees_path}}", priority="{{inputs.priority}}", deadline="{{inputs.deadline}}", effort_hours="{{inputs.effort_hours}}"), "requirements")
step("WF009", 2, "Compare employee skills", "match_employee_skills", dict(employees_path="{{inputs.employees_path}}", requirements="{{ctx.requirements}}"), "skill_match")
step("WF009", 3, "Check current workload", "check_workload", dict(rows="{{ctx.skill_match.rows}}", requirements="{{ctx.requirements}}"), "workload")
step("WF009", 4, "Rank candidates", "rank_candidates", dict(rows="{{ctx.workload.rows}}"), "ranked")
step("WF009", 5, "Select employee", "select_employee", dict(rows="{{ctx.ranked.rows}}", requirements="{{ctx.requirements}}"), "selection", note="Excel rule: prefer required skills + available capacity; ESCALATE if no suitable employee")
step("WF009", 6, "Generate assignment summary", "make_report", dict(title="Task assignment", summary="Recommended assignee: {{ctx.selection.selected.name}} ({{ctx.selection.selected.role}})", sections={"Reasoning": "{{ctx.selection.reasoning}}", "Priority": "{{ctx.requirements.priority}}", "Deadline": "{{ctx.requirements.deadline}} ({{ctx.requirements.deadline_source}})", "Estimated effort": "{{ctx.requirements.effort_hours}} hours ({{ctx.requirements.effort_source}})", "Required skills": "{{ctx.requirements.skills}}", "Task summary": "{{ctx.requirements.description}}"}, tables={"Candidate ranking (top 5)": {"rows": "{{ctx.selection.candidates}}"}}), "result")

# ---- WF010
step("WF010", 1, "Load execution logs", "load_table", dict(path="{{inputs.log_files}}", skip_missing=True, required_columns=["run_id", "workflow_id", "step_name", "status", "duration_sec"], optional_numeric_columns=["duration_sec"]), "logs", note="Includes logs/runtime_logs.csv written by this very agent")
step("WF010", 2, "Calculate success/failure rate", "calc_failure_rates", dict(rows="{{ctx.logs.rows}}"), "rates")
step("WF010", 3, "Calculate average execution time", "calc_execution_times", dict(rows="{{ctx.logs.rows}}"), "timings")
step("WF010", 4, "Identify frequent errors", "frequent_errors", dict(rows="{{ctx.logs.rows}}", top_n=5), "errors")
step("WF010", 5, "Identify slow steps", "slow_steps", dict(rows="{{ctx.logs.rows}}", threshold_sec="{{inputs.slow_step_threshold_sec}}", top_n=5), "slow")
step("WF010", 6, "Generate recommendations", "generate_recommendations", dict(rates="{{ctx.rates}}", timings="{{ctx.timings}}", errors="{{ctx.errors}}", slow="{{ctx.slow}}", failure_rate_threshold_pct="{{inputs.failure_rate_threshold_pct}}", exec_time_threshold_sec="{{inputs.exec_time_threshold_sec}}"), "recs", note="Excel rule: flag failure rate > 10% or avg time above threshold")
step("WF010", 7, "Generate performance report", "make_report", dict(title="Workflow performance report", summary="{{ctx.recs.flagged_count}} of {{ctx.rates.count}} workflows are flagged (failure rate > {{inputs.failure_rate_threshold_pct}}% or average time > {{inputs.exec_time_threshold_sec}}s), based on {{ctx.rates.total_runs}} runs.", tables={"Failure rate by workflow (worst first)": {"rows": "{{ctx.rates.rows}}", "columns": ["workflow_id", "workflow", "runs", "failed", "failure_rate_pct"], "export": True}, "Flagged workflows": {"rows": "{{ctx.recs.flagged}}", "columns": ["workflow_id", "workflow", "failure_rate_pct", "avg_exec_sec", "flags"]}, "Average execution time": {"rows": "{{ctx.timings.rows}}", "limit": 10}, "Most frequent errors": {"rows": "{{ctx.errors.rows}}"}, "Slow steps": {"rows": "{{ctx.slow.rows}}"}}, sections={"Recommendations": "{{ctx.recs.recommendations}}"}), "result")


# ---------------------------------------------------------------- write workbook
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY = Font(name="Arial", size=9)


def add_sheet(wb, title, header, rows, widths):
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = BODY, Alignment(vertical="top", wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


wb = openpyxl.load_workbook(SRC)
add_sheet(wb, "Input_Spec", ["Workflow_ID", "Input_Name", "Type", "Required", "Default", "Description", "Ask_Prompt", "Extract_Regex", "Fields"],
          INPUTS, [11, 24, 12, 9, 34, 46, 54, 52, 40])
add_sheet(wb, "Step_Bindings", ["Workflow_ID", "Step_No", "Step_Name", "Tool", "Args_JSON", "Output_Key", "Condition", "On_Error", "Note"],
          S, [11, 8, 38, 26, 90, 16, 34, 14, 56])
add_sheet(wb, "Routing_Hints", ["Workflow_ID", "Keywords"], HINTS.items(), [12, 110])
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT} sheets={wb.sheetnames} steps={len(S)} inputs={len(INPUTS)}")
